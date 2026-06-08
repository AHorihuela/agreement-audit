"""[/agreement-audit — step 1] Parse the agreements to text + resolve the audit scope.

LOCAL only (docx via python-docx, pdf via PyMuPDF, md/txt direct) — NO API call, so it consumes no
credits and no subscription usage. Writes one `<doc_id>.md` per parsed agreement under --out (default
.audit/, gitignored — it holds your contract text) and prints the JSON manifest the workflow fans out
over: {"doc_ids": [...], "skipped": [...], "fields": [...], "out": "<dir>"}.

Scope (which docs + which clauses) is resolved here, not via the workflow's `args` global, so a run is
fully controllable + reproducible. Precedence: --config file > --docs flag > default. The skill writes
`.audit/audit_config.json`; this reads it.

Usage:
  python3 scripts/audit_prep.py --config .audit/audit_config.json --out .audit
  python3 scripts/audit_prep.py --docs agreements --out .audit          # no config -> default clauses
"""
import argparse
import json
from pathlib import Path

DEFAULT_DOCS = "agreements"
SUPPORTED = {".docx", ".pdf", ".md", ".txt"}
DEFAULT_FIELDS = [
    {"key": "governing_law", "label": "Governing law",
     "question": "What is the governing law (choice of law) of this agreement?"},
    {"key": "exclusivity", "label": "Exclusivity",
     "question": "Are the granted rights exclusive or non-exclusive?"},
    {"key": "term", "label": "Term", "question": "What is the term / duration of the agreement?"},
    {"key": "termination_notice", "label": "Termination notice",
     "question": "What notice is required to terminate the agreement (for convenience)?"},
    {"key": "mfn", "label": "MFN",
     "question": "Is there a most-favored-nation / most-favored-customer clause?"},
]


def extract_text(path: Path) -> str:
    """Best-effort text extraction. Returns "" if unsupported or the file has no extractable text
    (e.g. a scanned PDF needing OCR) — the caller records that as a skip, never a silent empty."""
    ext = path.suffix.lower()
    if ext in (".md", ".txt"):
        return path.read_text(encoding="utf-8", errors="replace")
    if ext == ".docx":
        from docx import Document
        d = Document(str(path))
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for table in d.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
        return "\n".join(parts)
    if ext == ".pdf":
        import fitz  # PyMuPDF
        with fitz.open(str(path)) as doc:
            return "\n".join(page.get_text() for page in doc)
    return ""


def select_files(docs_dir, files) -> list:
    """Resolve the input files. Paths may be ANYWHERE on disk — absolute, `~/...`, or relative to the
    current directory; the documents do NOT have to live inside this repo. This is what lets the audit
    run on ONE agreement, a named subset, or a whole folder:
      - an explicit `files` list (each a path) -> exactly those;
      - `docs_dir` pointing at a single file -> just that file;
      - `docs_dir` pointing at a directory -> every file in it.
    """
    if files:
        return [Path(p).expanduser() for p in files]
    p = Path(docs_dir).expanduser()       # expand ~ so external/home paths resolve
    if p.is_dir():
        return [c for c in sorted(p.iterdir()) if c.is_file()]
    if p.is_file() or p.suffix:           # a real file, or a file-like path (typo) -> surface "not found"
        return [p]
    return []                             # a missing directory -> empty (the skill flags a zero-doc run)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--config", help="JSON {docs_dir|files, fields:[{key,label,question}]} — resolves scope")
    ap.add_argument("--docs", help="a directory OR a single file (overridden by config.docs_dir/files)")
    ap.add_argument("--out", default=".audit", help="where to write per-doc text (gitignore this)")
    args = ap.parse_args()

    cfg = {}
    if args.config and Path(args.config).is_file():
        cfg = json.loads(Path(args.config).read_text(encoding="utf-8"))
    docs_dir = cfg.get("docs_dir") or args.docs or DEFAULT_DOCS
    files_cfg = cfg.get("files")                          # optional explicit list of file paths
    fields = cfg.get("fields") or DEFAULT_FIELDS

    candidates = select_files(docs_dir, files_cfg)
    # An explicitly named file / list means the user picked these on purpose — surface a bad pick as a
    # reasoned skip rather than silently ignoring it (as we do for stray files in a scanned directory).
    dd = Path(docs_dir).expanduser()
    explicit = bool(files_cfg) or dd.is_file() or (bool(dd.suffix) and not dd.is_dir())

    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    doc_ids, skipped, seen = [], [], {}
    for path in candidates:
        if not path.is_file():
            if explicit:
                skipped.append({"doc_id": path.stem, "reason": "file not found"})
            continue
        if path.suffix.lower() not in SUPPORTED:
            if explicit:
                skipped.append({"doc_id": path.stem,
                                "reason": f"unsupported file type ({path.suffix or 'no extension'})"})
            continue
        doc_id = path.stem
        # Same-stem files (acme.docx + acme.pdf) would overwrite each other's .md and double-count the
        # coverage receipt — a silent loss of a document. Refuse rather than corrupt the audit.
        if doc_id in seen:
            raise SystemExit(f"Duplicate document id {doc_id!r}: '{seen[doc_id].name}' and "
                             f"'{path.name}' share a filename stem. Rename one — same-stem inputs "
                             f"would overwrite each other and under-count coverage.")
        seen[doc_id] = path
        try:
            text = extract_text(path).strip()
            reason = "" if text else "no extractable text (likely scanned/encrypted — needs OCR)"
        except Exception as e:                               # a real parse failure, not a scanned doc
            text, reason = "", f"parse error: {type(e).__name__}: {e}"
        if not text:                                         # record WHY, never a silent vanish
            skipped.append({"doc_id": doc_id, "reason": reason})
            continue
        (out / f"{doc_id}.md").write_text(text, encoding="utf-8")
        doc_ids.append(doc_id)
    print(json.dumps({"doc_ids": doc_ids, "skipped": skipped, "fields": fields, "out": str(out)}))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
