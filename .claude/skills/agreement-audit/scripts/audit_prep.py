"""[/agreement-audit — step 1] Parse the agreements to text + resolve the audit scope.

LOCAL only (docx via python-docx, pdf via PyMuPDF, md/txt direct) — NO API call, so it consumes no
credits and no subscription usage. Writes one `<doc_id>.md` per parsed agreement under `<out>/sources/`
(default `.audit/sources/`, gitignored — it holds your contract text) and prints the JSON manifest the
workflow fans out over: {"doc_ids": [...], "skipped": [{"doc_id","reason"}], "fields": [...], "sources": "<dir>"}.

Scope (which docs + which clauses) is resolved here, not via the workflow's `args` global, so a run is
fully controllable + reproducible. Precedence: --config file > --docs flag > default.

Usage:
  python3 scripts/audit_prep.py --config .audit/audit_config.json --out .audit
  python3 scripts/audit_prep.py --docs agreements --out .audit          # no config -> default clauses
"""
import argparse
import glob as _glob
import json
import re
from pathlib import Path

DEFAULT_DOCS = "agreements"
SUPPORTED = {".docx", ".pdf", ".md", ".txt"}
DEFAULT_FIELDS = [
    {"key": "governing_law", "label": "Governing law",
     "question": "What is the governing law (choice of law) of this agreement?"},
    {"key": "exclusivity", "label": "Exclusivity",
     "question": "Are the granted rights exclusive or non-exclusive?"},
    {"key": "term", "label": "Term", "question": "What is the term / duration of the agreement?"},
    {"key": "termination_notice", "label": "Termination notice",
     "question": "What notice is required to terminate the agreement (for convenience)?"},
    {"key": "mfn", "label": "MFN",
     "question": "Is there a most-favored-nation / most-favored-customer clause?"},
]


def safe_doc_id(stem: str) -> str:
    """A filesystem- AND shell-safe id from a filename stem. The per-doc text is written to
    `<sources>/<id>.md`, and the workflow agents `cat` that path inside a quoted shell command — so the
    id must not contain shell metacharacters (`$`, backtick, `"`, `;`, `|`, `&`, `<`, `>`, `(`, `)`,
    `\\`, …): an oddly- or maliciously-named file (`a$(cmd).pdf`) must not be able to inject or break
    the command. Normal names (letters, digits, spaces, dots, dashes, underscores) pass through."""
    s = re.sub(r"[^A-Za-z0-9._ -]", "_", stem).strip()
    return s or "document"


def extract_text(path: Path) -> str:
    """Best-effort text extraction. Returns "" if the file has no extractable text (e.g. a scanned PDF
    needing OCR) — the caller records that as a reasoned skip, never a silent empty."""
    ext = path.suffix.lower()
    if ext in (".md", ".txt"):
        return path.read_text(encoding="utf-8", errors="replace")
    if ext == ".docx":
        from docx import Document
        d = Document(str(path))
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for table in d.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
        return "\n".join(parts)
    if ext == ".pdf":
        import fitz  # PyMuPDF
        with fitz.open(str(path)) as doc:
            return "\n".join(page.get_text() for page in doc)
    return ""


def select_files(docs_dir, files) -> list:
    """Resolve the input files. Paths may be ANYWHERE on disk — absolute, `~/...`, a glob, or relative
    to the current directory; the documents do NOT have to live inside this repo. This is what lets the
    audit run on ONE agreement, a named subset, a glob, or a whole folder:
      - an explicit `files` list (each a path) -> exactly those;
      - `docs_dir` pointing at a real file -> just that file (even if its name contains `[`/`]`);
      - `docs_dir` pointing at a directory -> every file in it;
      - otherwise `docs_dir` containing glob characters (`*?[`) -> the matches.

    A LITERAL path is checked before glob interpretation, so a real filename that happens to contain
    glob metacharacters — e.g. `LC_Redline ... [vs Current].pdf` — resolves to itself instead of being
    parsed as a wildcard character-class (which matches a single char, and so would match nothing).
    """
    if files:
        return [Path(p).expanduser() for p in files]
    raw = str(Path(docs_dir).expanduser())                # expand ~ so external/home paths resolve
    p = Path(raw)
    if p.is_file():                                       # a real file wins over glob interpretation
        return [p]
    if p.is_dir():
        return [c for c in sorted(p.iterdir()) if c.is_file()]
    if any(ch in raw for ch in "*?["):                    # not a literal path -> treat as a glob
        matches = sorted(_glob.glob(raw))
        if matches:
            return [Path(m) for m in matches]
        # No matches: a bracketed name with no `*`/`?` is almost always a literal file the user typo'd,
        # not an empty wildcard — surface it as not-found rather than vanishing it. A true `*`/`?` glob
        # that matched nothing stays an empty (0-doc) run.
        if p.suffix and not any(ch in raw for ch in "*?"):
            return [p]
        return []
    if p.suffix:                                          # a file-like path that doesn't exist (typo)
        return [p]
    return []                                             # a missing directory -> empty (skill flags 0-doc)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--config", help="JSON {docs_dir|files, fields:[{key,label,question}]} — resolves scope")
    ap.add_argument("--docs", help="a directory, a single file, or a glob (overridden by config.docs_dir/files)")
    ap.add_argument("--out", default=".audit", help="working dir; per-doc text goes to <out>/sources/")
    args = ap.parse_args()

    cfg = {}
    if args.config and Path(args.config).is_file():
        cfg = json.loads(Path(args.config).read_text(encoding="utf-8"))
    docs_dir = cfg.get("docs_dir") or args.docs or DEFAULT_DOCS
    files_cfg = cfg.get("files")                          # optional explicit list of file paths
    fields = cfg.get("fields") or DEFAULT_FIELDS

    candidates = select_files(docs_dir, files_cfg)
    # An explicitly named file / list / glob means the user picked these on purpose — surface a bad pick
    # as a reasoned skip rather than silently ignoring it (as we do for OS noise in a scanned directory).
    dd = Path(docs_dir).expanduser()
    has_glob = files_cfg is None and any(ch in str(dd) for ch in "*?[")
    explicit = bool(files_cfg) or has_glob or dd.is_file() or (bool(dd.suffix) and not dd.is_dir())

    # Per-doc sources live in their own subdir, wiped each run, so a stale source from a PRIOR run can
    # never be grounded against (and can't collide with the report file).
    src_dir = Path(args.out) / "sources"
    src_dir.mkdir(parents=True, exist_ok=True)
    for stale in src_dir.glob("*.md"):
        stale.unlink()

    doc_ids, skipped, seen = [], [], {}
    for path in candidates:
        if path.is_dir():                                # a directory handed to the `files` list
            skipped.append({"doc_id": path.name, "reason": "is a directory — pass it as the folder, "
                            "or list its files individually"})
            continue
        if not path.exists():
            if explicit:
                skipped.append({"doc_id": path.stem or path.name, "reason": "path not found"})
            continue
        if path.suffix.lower() not in SUPPORTED:
            # Surface a real-looking unsupported file (an explicit pick, OR a non-dotfile in a scanned
            # directory) so a wrong-type or extensionless agreement is visible, never silently dropped.
            if explicit or not path.name.startswith("."):
                skipped.append({"doc_id": path.stem or path.name,
                                "reason": f"unsupported file type ({path.suffix or 'no extension'}) — "
                                          f"convert to .pdf/.docx/.md/.txt"})
            continue
        doc_id = safe_doc_id(path.stem)
        # Same id (after sanitizing) would overwrite the prior source + double-count the receipt — a
        # silent loss of a document. Refuse rather than corrupt the audit.
        if doc_id in seen:
            raise SystemExit(f"Duplicate document id {doc_id!r}: '{seen[doc_id].name}' and "
                             f"'{path.name}' map to the same id. Rename one — they would overwrite "
                             f"each other and under-count coverage.")
        seen[doc_id] = path
        try:
            text = extract_text(path).strip()
            reason = "" if text else "no extractable text (likely scanned/encrypted — needs OCR)"
        except Exception as e:                               # a real parse failure, not a scanned doc
            text, reason = "", f"parse error: {type(e).__name__}: {e}"
        if not text:                                         # record WHY, never a silent vanish
            skipped.append({"doc_id": doc_id, "reason": reason})
            continue
        (src_dir / f"{doc_id}.md").write_text(text, encoding="utf-8")
        doc_ids.append(doc_id)
    manifest = {"doc_ids": doc_ids, "skipped": skipped, "fields": fields, "sources": str(src_dir)}
    # Persist the manifest so the (script-free) workflow can read scope from it, and the skill can
    # reuse it to assemble findings.json — one source of truth, no reliance on the workflow `args`.
    (Path(args.out) / "manifest.json").write_text(json.dumps(manifest), encoding="utf-8")
    print(json.dumps(manifest))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
